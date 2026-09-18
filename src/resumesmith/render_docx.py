"""Word export: a real, editable .docx laid out like the PDF (same theme, same fit scale)."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from .model import Resume
from .text import contact_lines, date_range, full_url, group_jobs, segments, visible_sections

# Fonts every Word install has, closest to each theme's PDF font.
LOOKS = {
    "classic": {"font": "Georgia", "size": 10.0, "center": True, "italic_headline": True, "accent": False},
    "modern": {"font": "Calibri", "size": 10.5, "center": False, "italic_headline": False, "accent": True},
    "compact": {"font": "Arial", "size": 9.5, "center": False, "italic_headline": False, "accent": True},
}
MUTED = RGBColor(0x44, 0x44, 0x44)
INK = RGBColor(0x16, 0x16, 0x16)

# Elements that must come after <w:pBdr> inside <w:pPr> (Word rejects out-of-order XML).
_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
               "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
               "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
               "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
               "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


class _Writer:
    def __init__(self, r: Resume, scale: float, margin_mm: float):
        self.r = r
        self.look = LOOKS.get(r.settings.theme, LOOKS["classic"])
        self.size = self.look["size"] * scale
        self.accent = RGBColor.from_string(r.settings.accent.lstrip("#").upper())
        self.doc = Document()

        sec = self.doc.sections[0]
        if r.settings.page == "A4":
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(margin_mm)
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin

        normal = self.doc.styles["Normal"]
        normal.font.name = self.look["font"]
        normal.font.size = Pt(self.size)
        normal.font.color.rgb = INK
        rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), self.look["font"])
        pf = normal.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.08

        props = self.doc.core_properties
        props.author = r.basics.name
        props.title = f"{r.basics.name} – Resume"

    # ---------- primitives ----------

    def para(self, before: float = 0, after: float = 0, keep_next: bool = False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = keep_next
        return p

    def run(self, p, text: str, bold=False, italic=False, color=None, size=None, href=None):
        run = p.add_run(text)
        run.bold, run.italic = bold or None, italic or None
        if color is not None:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
        if href:
            rel = p.part.relate_to(href, RT.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rel)
            p._p.remove(run._r)
            link.append(run._r)
            p._p.append(link)
        return run

    def rich(self, p, text: str, **fmt):
        """Text with **bold** and [links](url)."""
        for seg in segments(text):
            self.run(p, seg.text, **{**fmt, "bold": fmt.get("bold") or seg.bold, "href": seg.href})

    def row(self, left: list[tuple[str, dict]], right: str | None, before: float = 0):
        """Left-aligned text with a right-aligned date/location, like the PDF's two-column rows."""
        p = self.para(before=before, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        for text, fmt in left:
            self.run(p, text, **fmt)
        if right:
            self.run(p, "\t" + right, color=MUTED, size=self.size * 0.95)
        return p

    def bullet(self, text: str):
        p = self.doc.add_paragraph(style="List Bullet")
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Pt(self.size * 1.5), Pt(-self.size * 1.0)
        pf.space_before, pf.space_after = Pt(self.size * 0.15), Pt(0)
        self.rich(p, text)

    def heading(self, title: str):
        p = self.para(before=self.size * 0.95, after=self.size * 0.4, keep_next=True)
        color = self.accent if self.look["accent"] else INK
        self.run(p, title.upper(), bold=True, color=color, size=self.size)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": str(color)}.items():
            bottom.set(qn(k), v)
        border.append(bottom)
        p._p.get_or_add_pPr().insert_element_before(border, *_AFTER_PBDR)

    # ---------- sections ----------

    def header(self):
        b = self.r.basics
        align = WD_ALIGN_PARAGRAPH.CENTER if self.look["center"] else WD_ALIGN_PARAGRAPH.LEFT
        p = self.para()
        p.alignment = align
        self.run(p, b.name, bold=True, size=self.size * 2.2,
                 color=self.accent if self.look["accent"] else INK)
        if b.title:
            p = self.para(before=self.size * 0.2)
            p.alignment = align
            self.run(p, b.title, italic=self.look["italic_headline"], size=self.size * 1.1)
        for n, line in enumerate(contact_lines(self.r)):
            p = self.para(before=self.size * (0.1 if n else 0.35))
            p.alignment = align
            for i, (text, href) in enumerate(line):
                if i:
                    self.run(p, "  ·  ", color=MUTED, size=self.size * 0.93)
                self.run(p, text, color=MUTED, size=self.size * 0.93, href=href)

    def summary(self):
        self.rich(self.para(), self.r.summary)

    def skills(self):
        for i, g in enumerate(self.r.skills):
            p = self.para(before=self.size * 0.12 if i else 0)
            self.run(p, f"{g.category}: ", bold=True)
            self.run(p, ", ".join(g.items))

    def experience(self):
        for i, g in enumerate(group_jobs(self.r.experience)):
            self.row([(g.company, {"bold": True, "href": full_url(g.url) if g.url else None})], g.location,
                     before=self.size * 0.55 if i else 0)
            for k, j in enumerate(g.jobs):
                self.row([(j.role, {"italic": True})], date_range(j.start, j.end),
                         before=self.size * 0.3 if k else 0)
                if j.summary:
                    self.rich(self.para(), j.summary, italic=True, color=MUTED)
                for b in j.bullets:
                    self.bullet(b)
                if j.tech:
                    p = self.para(before=self.size * 0.15)
                    self.run(p, "Tech: ", bold=True, size=self.size * 0.93)
                    self.run(p, ", ".join(j.tech), color=MUTED, size=self.size * 0.93)

    def projects(self):
        for i, pr in enumerate(self.r.projects):
            left = [(pr.name, {"bold": True, "href": full_url(pr.url) if pr.url else None})]
            if pr.tech:
                left.append((" | " + ", ".join(pr.tech), {"italic": True, "color": MUTED}))
            dates = date_range(pr.start, pr.end) if (pr.start or pr.end) else None
            self.row(left, dates, before=self.size * 0.5 if i else 0)
            if pr.description:
                self.rich(self.para(), pr.description, italic=True, color=MUTED)
            for b in pr.bullets:
                self.bullet(b)

    def education(self):
        for i, e in enumerate(self.r.education):
            self.row([(e.school, {"bold": True})], e.location, before=self.size * 0.5 if i else 0)
            left = [(e.degree, {"italic": True})]
            if e.score:
                left.append((" — " + e.score, {"italic": True}))
            self.row(left, date_range(e.start, e.end) if (e.start or e.end) else None)
            for b in e.bullets:
                self.bullet(b)

    def certifications(self):
        for c in self.r.certifications:
            text = f"[{c.name}]({c.url})" if c.url else c.name
            self.bullet(text + (f" — {c.issuer}" if c.issuer else "") + (f" ({date_range(c.date, None)})" if c.date else ""))

    def build(self) -> bytes:
        self.header()
        for key, title in visible_sections(self.r):
            self.heading(title)
            if key == "achievements":
                for a in self.r.achievements:
                    self.bullet(a)
            elif key.startswith("extra:"):
                for item in self.r.extra[int(key.split(":")[1])].items:
                    self.bullet(item)
            else:
                getattr(self, key)()
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


def render_docx(r: Resume, scale: float = 1.0, margin_mm: float | None = None) -> bytes:
    return _Writer(r, scale, margin_mm if margin_mm is not None else r.settings.margin_mm).build()
